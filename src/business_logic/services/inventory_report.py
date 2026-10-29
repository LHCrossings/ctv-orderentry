"""Inventory (avails) report — how much commercial break time is still unsold, per program.

2026-10-29, Maija/Lee for the AEs (Charmaine): for a date range, optionally a broadcast time range,
any set of markets and languages, show per program block the COMS break time that is still open
("what has not been filled that can be sold") and how many breaks still have room — a show with
four breaks will never place five spots, however much time is open. Maija's hand-made sheet is the
row shape: LANGUAGE | PROGRAM DAY/TIME | TITLE | BREAKS; this adds the open time over the range and
a day-by-day detail, and exports to Excel (print the page for a PDF).

Data (all from Etere, read-only):
  * capacity — the grid: Traffic_Calendar (Level 0 = published) → traffic_scheduleblock → traffic_block
    (one per market, name "<lang prefix> - <title>") → traffic_segment of Type 'COMS' (visible). A
    segment's absolute offset is block offset + segment offset, frame-of-day on the 06:00→30:00
    broadcast scale; Duration is the break length in frames.
  * sold — placed contract spots: trafficPalinse (the contract side) joined to TPALINSE (the playlist),
    LIVELLO 0, NEWTYPE COM/COMS/BNS/AV. `trafficPalinse.offset` IS the break the spot was booked into
    (the block column EE shows) and equals a COMS segment's absolute offset for 98% of rows; the rest
    are attributed to the nearest break of the block their ORA falls in. PI/PSA/promo/ID fillers never
    carry trafficPalinse and stay "open" — they are what gets bumped when a spot sells.
  * language — the block-name prefix (M, C, V, K, T, P, SA, Hm, J, E), same letters the block tab uses.

A break "has room" when its open time is at least MIN_SELLABLE_S (one :15).
"""

from __future__ import annotations

import datetime as dt
import re
from collections import defaultdict

FPS = 29.97
HOUR_FRAMES = 107892  # 3600 * 29.97
DAY_FRAMES = 24 * HOUR_FRAMES
MIN_SELLABLE_S = 15
SOLD_TYPES = ("COM", "COMS", "BNS", "AV")

MARKETS = {
    1: "NYC",
    2: "CMP",
    3: "HOU",
    4: "SFO",
    5: "SEA",
    6: "LAX",
    7: "CVC",
    8: "WDC",
    9: "MMT",
    10: "DAL",
}
LANGUAGES = {
    "M": "Mandarin",
    "C": "Cantonese",
    "V": "Vietnamese",
    "K": "Korean",
    "T": "Filipino/Tagalog",
    "P": "Punjabi",
    "SA": "South Asian",
    "Hm": "Hmong",
    "J": "Japanese",
    "E": "English",
}
_DAY_LETTERS = ["M", "Tu", "W", "Th", "F", "Sa", "Su"]
_NAME_RE = re.compile(r"^\s*([A-Za-z]{1,2})\s*-\s*(.*?)\s*$")


def parse_block_name(name: str) -> tuple[str, str]:
    """'P - To the Point REV' -> ('P', 'To the Point REV'); 'K- Korean Drama Saturday' -> ('K', ...).
    A name without a recognised prefix is language '?' with the whole name as title."""
    m = _NAME_RE.match(name or "")
    if m:
        pfx = m.group(1)
        for known in LANGUAGES:
            if pfx.lower() == known.lower():
                return known, m.group(2)
    return "?", (name or "").strip()


def frames_to_bcast(frames: int) -> str:
    """Frame-of-day -> '2p', '3:30p', '12a' (post-midnight tail wraps past 24h)."""
    total_min = int(round(frames / HOUR_FRAMES * 60))
    h, m = (total_min // 60) % 24, total_min % 60
    suffix = "a" if h < 12 else "p"
    h12 = h % 12 or 12
    return f"{h12}{suffix}" if m == 0 else f"{h12}:{m:02d}{suffix}"


def bcast_frames(hhmm: str | None) -> int | None:
    """'14:00' -> frames; hours < 6 are the post-midnight tail (+24h). None/blank -> None."""
    m = re.match(r"\s*(\d{1,2}):(\d{2})", hhmm or "")
    if not m:
        return None
    h, mins = int(m.group(1)), int(m.group(2))
    if h < 6:
        h += 24
    return int(round((h * 60 + mins) * 60 * FPS))


def day_pattern(weekdays: set[int]) -> str:
    """Set of Monday=0..Sunday=6 -> 'M-F', 'Sa', 'Su', 'Sa-Su', 'M-Su', 'M-Sa', else 'M,W,F'."""
    s = set(weekdays)
    named = {
        frozenset(range(7)): "M-Su",
        frozenset(range(6)): "M-Sa",
        frozenset(range(5)): "M-F",
        frozenset({5, 6}): "Sa-Su",
        frozenset({5}): "Sa",
        frozenset({6}): "Su",
    }
    if frozenset(s) in named:
        return named[frozenset(s)]
    return ",".join(_DAY_LETTERS[d] for d in sorted(s)) or "-"


# ---------------------------------------------------------------------------
# Fetch (read-only)
# ---------------------------------------------------------------------------
def fetch_segments(conn, markets: list[int], d_from: dt.date, d_to: dt.date) -> list[dict]:
    """Every visible COMS segment of every published grid day in range for the markets."""
    if not markets:
        return []
    cur = conn.cursor()
    ph = ",".join("%s" for _ in markets)
    cur.execute(
        f"""SELECT tc.Date, tb.cod_user, tb.id_trafficblock, RTRIM(tb.name), ts.Offset, tb.Duration,
                   ts.Offset + tseg.Offset, tseg.Duration, tseg.ID_TrafficSegment
            FROM Traffic_Calendar tc
            JOIN traffic_scheduleblock ts ON ts.id_trafficschedule = tc.id_trafficschedule
            JOIN traffic_block tb ON tb.id_trafficblock = ts.id_trafficblock
            JOIN traffic_segment tseg ON tseg.ID_TrafficBlock = tb.id_trafficblock
            WHERE tc.Date BETWEEN %s AND %s AND tc.Level = 0 AND tb.cod_user IN ({ph})
              AND tseg.visible = 1 AND tseg.Type = 'COMS' AND ISNULL(tb.expired, 0) = 0
            ORDER BY tc.Date, tb.cod_user, ts.Offset, tseg.Offset""",
        (d_from.isoformat(), d_to.isoformat(), *markets),
    )
    out = []
    for d, cu, bid, name, boff, bdur, soff, sdur, sid in cur.fetchall():
        out.append(
            {
                "date": d.date().isoformat() if hasattr(d, "date") else str(d)[:10],
                "market": int(cu),
                "block_id": int(bid),
                "name": name or "",
                "bstart": int(boff),
                "bdur": int(bdur or 0),
                "seg_off": int(soff),
                "seg_dur": int(sdur or 0),
                "seg_id": int(sid),
            }
        )
    return out


def fetch_sold(conn, markets: list[int], d_from: dt.date, d_to: dt.date) -> list[dict]:
    """Placed contract spots in range: (market, date, booked break offset, ORA, duration frames)."""
    if not markets:
        return []
    cur = conn.cursor()
    ph = ",".join("%s" for _ in markets)
    types = ",".join(f"'{t}'" for t in SOLD_TYPES)
    cur.execute(
        f"""SELECT t.DATA, t.COD_USER, tp.offset, t.ORA, t.DURATION
            FROM trafficPalinse tp JOIN TPALINSE t ON t.ID_TPALINSE = tp.id_tpalinse
            WHERE t.DATA BETWEEN %s AND %s AND t.COD_USER IN ({ph}) AND t.LIVELLO = 0
              AND t.NEWTYPE IN ({types})""",
        (d_from.isoformat(), d_to.isoformat(), *markets),
    )
    return [
        {
            "date": d.date().isoformat() if hasattr(d, "date") else str(d)[:10],
            "market": int(cu),
            "offset": int(off or 0),
            "ora": int(ora or 0),
            "dur": int(dur or 0),
        }
        for d, cu, off, ora, dur in cur.fetchall()
    ]


# ---------------------------------------------------------------------------
# Build (pure)
# ---------------------------------------------------------------------------
def _secs(frames: int) -> int:
    return int(round(frames / FPS))


def build(
    segments: list[dict],
    sold: list[dict],
    *,
    languages: list[str] | None = None,
    t_from: str | None = None,
    t_to: str | None = None,
    min_sellable_s: int = MIN_SELLABLE_S,
) -> dict:
    """Aggregate segments + sold spots into the report: per market, rows per program block
    (name + start time) with the day pattern, breaks per airing, airings in range, capacity,
    sold, open seconds, open per airing, breaks with room; plus per-date detail per row."""
    lo, hi = bcast_frames(t_from), bcast_frames(t_to)
    if lo is not None and hi is not None and hi <= lo:
        hi += DAY_FRAMES
    langs = {x for x in (languages or [])}

    # block-days: (market, date, block_id) -> block with its segments
    blocks: dict[tuple, dict] = {}
    for s in segments:
        lang, title = parse_block_name(s["name"])
        if langs and lang not in langs:
            continue
        if lo is not None and s["bstart"] < lo:
            continue
        if hi is not None and s["bstart"] >= hi:
            continue
        key = (s["market"], s["date"], s["block_id"], s["bstart"])
        b = blocks.setdefault(
            key,
            {
                "market": s["market"],
                "date": s["date"],
                "block_id": s["block_id"],
                "name": s["name"],
                "lang": lang,
                "title": title,
                "bstart": s["bstart"],
                "bdur": s["bdur"],
                "segs": {},
            },
        )
        b["segs"].setdefault(s["seg_off"], {"cap": 0, "sold": 0, "spots": 0})
        b["segs"][s["seg_off"]]["cap"] += s["seg_dur"]

    # index for attribution
    by_md: dict[tuple, list[dict]] = defaultdict(list)
    for b in blocks.values():
        by_md[(b["market"], b["date"])].append(b)
    unattributed = 0
    for sp in sold:
        cands = by_md.get((sp["market"], sp["date"]), [])
        hit = None
        for b in cands:
            if sp["offset"] in b["segs"]:
                hit = (b, sp["offset"])
                break
        if hit is None:
            for b in cands:
                if b["bstart"] <= sp["ora"] < b["bstart"] + max(b["bdur"], 1) and b["segs"]:
                    off = min(b["segs"], key=lambda o: abs(o - sp["ora"]))
                    hit = (b, off)
                    break
        if hit is None:
            unattributed += 1  # a spot in a block outside the filters, or no grid for that day
            continue
        seg = hit[0]["segs"][hit[1]]
        seg["sold"] += sp["dur"]
        seg["spots"] += 1

    # per block-day numbers
    for b in blocks.values():
        cap = sum(x["cap"] for x in b["segs"].values())
        sld = sum(x["sold"] for x in b["segs"].values())
        b["cap_s"] = _secs(cap)
        b["sold_s"] = _secs(sld)
        b["open_s"] = max(0, _secs(cap - sld))
        b["breaks"] = len(b["segs"])
        b["breaks_open"] = sum(
            1 for x in b["segs"].values() if _secs(x["cap"] - x["sold"]) >= min_sellable_s
        )
        b["spots"] = sum(x["spots"] for x in b["segs"].values())

    dates = sorted({b["date"] for b in blocks.values()})
    # rows: (market, name, bstart) over the range
    rows: dict[tuple, dict] = {}
    for b in sorted(
        blocks.values(), key=lambda x: (x["market"], x["bstart"], x["name"], x["date"])
    ):
        k = (b["market"], b["name"], b["bstart"])
        r = rows.setdefault(
            k,
            {
                "market": b["market"],
                "market_code": MARKETS.get(b["market"], str(b["market"])),
                "lang": b["lang"],
                "lang_name": LANGUAGES.get(b["lang"], "Other"),
                "title": b["title"],
                "name": b["name"],
                "bstart": b["bstart"],
                "time": f"{frames_to_bcast(b['bstart'])}-{frames_to_bcast(b['bstart'] + b['bdur'])}",
                "weekdays": set(),
                "airings": 0,
                "breaks_per_airing": set(),
                "cap_s": 0,
                "sold_s": 0,
                "open_s": 0,
                "breaks": 0,
                "breaks_open": 0,
                "spots": 0,
                "days": {},
            },
        )
        wd = dt.date.fromisoformat(b["date"]).weekday()
        r["weekdays"].add(wd)
        r["airings"] += 1
        r["breaks_per_airing"].add(b["breaks"])
        for f in ("cap_s", "sold_s", "open_s", "breaks", "breaks_open", "spots"):
            r[f] += b[f]
        r["days"][b["date"]] = {
            "cap_s": b["cap_s"],
            "sold_s": b["sold_s"],
            "open_s": b["open_s"],
            "breaks": b["breaks"],
            "breaks_open": b["breaks_open"],
            "spots": b["spots"],
        }
    out_rows = []
    for r in rows.values():
        bpa = sorted(r["breaks_per_airing"])
        r["day_pattern"] = day_pattern(r["weekdays"])
        r["breaks_per_airing"] = bpa[0] if len(bpa) == 1 else f"{bpa[0]}-{bpa[-1]}"
        r["open_per_airing_s"] = round(r["open_s"] / r["airings"]) if r["airings"] else 0
        r["cap_per_airing_s"] = round(r["cap_s"] / r["airings"]) if r["airings"] else 0
        r.pop("weekdays")
        out_rows.append(r)

    markets_out = []
    for mid in sorted({r["market"] for r in out_rows}):
        mrows = [r for r in out_rows if r["market"] == mid]
        groups = []
        # language groups in broadcast-day order (Maija's sheet: Punjabi 2p, then Filipino 4p)
        first_air = {}
        for r in mrows:
            first_air[r["lang"]] = min(first_air.get(r["lang"], r["bstart"]), r["bstart"])
        for lang in sorted(first_air, key=lambda x: first_air[x]):
            lrows = sorted(
                (r for r in mrows if r["lang"] == lang),
                key=lambda r: (r["bstart"], r["day_pattern"]),
            )
            groups.append(
                {
                    "lang": lang,
                    "lang_name": LANGUAGES.get(lang, "Other"),
                    "rows": lrows,
                    "open_s": sum(r["open_s"] for r in lrows),
                    "breaks_open": sum(r["breaks_open"] for r in lrows),
                }
            )
        markets_out.append(
            {
                "market": mid,
                "market_code": MARKETS.get(mid, str(mid)),
                "groups": groups,
                "cap_s": sum(r["cap_s"] for r in mrows),
                "sold_s": sum(r["sold_s"] for r in mrows),
                "open_s": sum(r["open_s"] for r in mrows),
                "breaks": sum(r["breaks"] for r in mrows),
                "breaks_open": sum(r["breaks_open"] for r in mrows),
            }
        )
    return {
        "dates": dates,
        "markets": markets_out,
        "totals": {
            "cap_s": sum(m["cap_s"] for m in markets_out),
            "sold_s": sum(m["sold_s"] for m in markets_out),
            "open_s": sum(m["open_s"] for m in markets_out),
            "breaks": sum(m["breaks"] for m in markets_out),
            "breaks_open": sum(m["breaks_open"] for m in markets_out),
            "unattributed_spots": unattributed,
        },
        "min_sellable_s": min_sellable_s,
    }


def scan(
    conn,
    d_from: dt.date,
    d_to: dt.date,
    markets: list[int],
    languages: list[str] | None = None,
    t_from: str | None = None,
    t_to: str | None = None,
) -> dict:
    segs = fetch_segments(conn, markets, d_from, d_to)
    sold = fetch_sold(conn, markets, d_from, d_to)
    rep = build(segs, sold, languages=languages, t_from=t_from, t_to=t_to)
    rep.update(
        {
            "date_from": d_from.isoformat(),
            "date_to": d_to.isoformat(),
            "time_from": t_from or "",
            "time_to": t_to or "",
            "market_ids": markets,
            "languages": languages or [],
            "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        }
    )
    return rep


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def build_excel(rep: dict) -> bytes:
    import io

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def fill(c):
        return PatternFill("solid", fgColor=c)

    thin = Side(style="thin", color="BFBFBF")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr = Alignment(horizontal="center", vertical="center")
    NAVY, LT, GREY = "1F3864", "D6E4F0", "F2F2F2"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    span = f"{rep['date_from']} to {rep['date_to']}"
    if rep.get("time_from") or rep.get("time_to"):
        span += f", {rep.get('time_from') or '06:00'}-{rep.get('time_to') or '05:59'}"
    ws["A1"] = f"Inventory Report — open commercial time, {span}"
    ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    ws["A2"] = (
        "Open = COMS break time not taken by a placed paid, bonus or added-value spot; fillers (PI/PSA/promo) count as open. "
        f"A break has room when at least {rep['min_sellable_s']} s is open. Generated {rep.get('generated_at', '')}."
    )
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    hdr = [
        "Market",
        "Language",
        "Program day/time",
        "Title",
        "Breaks",
        "Airings",
        "Capacity/airing (s)",
        "Sold (s)",
        "Open (s)",
        "Open/airing (s)",
        "Breaks with room",
        "Breaks total",
    ]
    r = 4
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = ctr
        c.border = bdr
    r += 1
    for m in rep["markets"]:
        for g in m["groups"]:
            for row in g["rows"]:
                vals = [
                    m["market_code"],
                    row["lang_name"],
                    f"{row['day_pattern']} {row['time']}",
                    row["title"],
                    row["breaks_per_airing"],
                    row["airings"],
                    row["cap_per_airing_s"],
                    row["sold_s"],
                    row["open_s"],
                    row["open_per_airing_s"],
                    row["breaks_open"],
                    row["breaks"],
                ]
                for i, v in enumerate(vals, 1):
                    c = ws.cell(row=r, column=i, value=v)
                    c.border = bdr
                    if i >= 5:
                        c.alignment = ctr
                r += 1
            c = ws.cell(row=r, column=1, value=f"{m['market_code']} {g['lang_name']} total")
            c.font = Font(bold=True)
            for i in range(1, len(hdr) + 1):
                ws.cell(row=r, column=i).fill = fill(LT)
                ws.cell(row=r, column=i).border = bdr
            ws.cell(row=r, column=9, value=g["open_s"]).font = Font(bold=True)
            ws.cell(row=r, column=11, value=g["breaks_open"]).font = Font(bold=True)
            ws.cell(row=r, column=9).alignment = ctr
            ws.cell(row=r, column=11).alignment = ctr
            r += 1
        r += 1
    widths = [8, 16, 20, 34, 8, 8, 18, 10, 10, 15, 16, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # one daily sheet per market: rows = program, columns = dates (cap / sold / open / breaks with room)
    dates = rep["dates"]
    for m in rep["markets"]:
        wsd = wb.create_sheet(f"{m['market_code']} daily"[:31])
        wsd["A1"] = f"{m['market_code']} — open seconds per program per day ({span})"
        wsd["A1"].font = Font(bold=True, size=12, color=NAVY)
        sub = ["Cap", "Sold", "Open", "Brk room"]
        wsd.cell(row=3, column=1, value="Language").font = Font(bold=True)
        wsd.cell(row=3, column=2, value="Program day/time").font = Font(bold=True)
        wsd.cell(row=3, column=3, value="Title").font = Font(bold=True)
        col = 4
        for d in dates:
            wsd.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 3)
            c = wsd.cell(row=3, column=col, value=dt.date.fromisoformat(d).strftime("%a %m/%d"))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill(NAVY)
            c.alignment = ctr
            for j, s in enumerate(sub):
                cc = wsd.cell(row=4, column=col + j, value=s)
                cc.font = Font(bold=True, size=9)
                cc.fill = fill(GREY)
                cc.alignment = ctr
                cc.border = bdr
            col += 4
        rr = 5
        for g in m["groups"]:
            for row in g["rows"]:
                wsd.cell(row=rr, column=1, value=row["lang_name"])
                wsd.cell(row=rr, column=2, value=f"{row['day_pattern']} {row['time']}")
                wsd.cell(row=rr, column=3, value=row["title"])
                col = 4
                for d in dates:
                    day = row["days"].get(d)
                    vals = (
                        [
                            day["cap_s"],
                            day["sold_s"],
                            day["open_s"],
                            f"{day['breaks_open']}/{day['breaks']}",
                        ]
                        if day
                        else ["", "", "", ""]
                    )
                    for j, v in enumerate(vals):
                        c = wsd.cell(row=rr, column=col + j, value=v)
                        c.alignment = ctr
                        c.border = bdr
                        if j == 2 and day:
                            c.font = Font(bold=True)
                    col += 4
                rr += 1
        wsd.column_dimensions["A"].width = 15
        wsd.column_dimensions["B"].width = 20
        wsd.column_dimensions["C"].width = 32
        for i in range(4, 4 + 4 * len(dates)):
            wsd.column_dimensions[get_column_letter(i)].width = 7
        wsd.freeze_panes = "D5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
