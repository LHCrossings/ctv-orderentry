"""Agency XML Export: proposal workbook → TVB SpotTV Cable Proposal XML.

Oracle: the BAAQMD 2026 R1C file Kurt and Charmaine produced by hand and the
agency ingested (tests/fixtures/crispin/baaqmd_2026_r1c_agency_ingested.xml),
built from the REV1 proposal workbook beside it.
"""

from __future__ import annotations

import io
import json
import sys
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest

_root = Path(__file__).resolve().parents[2]
if str(_root) not in sys.path:
    sys.path.insert(0, str(_root))

from browser_automation.generators import agency_xml_export as ax
from browser_automation.parsers.crispin_parser import parse_crispin_xlsx

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "crispin"
WORKBOOK = FIX / "proposal_baaqmd_2026_rev1.xlsm"
ORACLE = FIX / "baaqmd_2026_r1c_agency_ingested.xml"

KURT_HEADER = dict(
    proposal_id="BAAQMD-2026-R1C",
    proposal_name="Crossings TV - BAAQMD 2026",
    advertiser="Bay Area Quality Management District",
    product="BAAQMD 2026",
    buyer_company="Allison Worldwide",
    buyer_name="Alexander Boyle",
    call_letters="3131CA",
    send_date=date(2026, 7, 23),
)


def _tree(el):
    return (
        el.tag,
        tuple(sorted(el.attrib.items())),
        (el.text or "").strip(),
        [_tree(c) for c in el],
    )


@pytest.fixture(scope="module")
def order():
    return parse_crispin_xlsx(str(WORKBOOK))


def test_reproduces_the_file_the_agency_ingested(order):
    xml_bytes = ax.export_proposal(order, ax.ExportHeader(**KURT_HEADER))
    assert _tree(ET.fromstring(xml_bytes)) == _tree(ET.parse(ORACLE).getroot())
    assert b'xmlns:tvb="http://www.AAAA.org/schemas/spotTV"' in xml_bytes
    assert xml_bytes.startswith(b'<?xml version="1.0" encoding="UTF-8"?>\n')


def test_oracle_and_output_validate_against_the_xsd(order):
    assert ax.validate_proposal_xml(ORACLE.read_bytes()) == []
    xml_bytes = ax.build_proposal_xml(ax.build_preview(order, ax.ExportHeader(**KURT_HEADER)))
    assert ax.validate_proposal_xml(xml_bytes) == []


def test_default_header_comes_from_the_sheet(order):
    h = ax.default_header(order)
    assert h.proposal_id == "BAAQMD-2026"
    assert h.proposal_name == "Crossings TV - BAAQMD 2026"
    assert h.product == "BAAQMD 2026"
    assert h.advertiser == "Bay Area Quality Management District"
    assert h.buyer_company == "Crispin LLC"
    assert h.buyer_name == "Alexander Boyle"
    assert h.call_letters == ""  # the buyer's station code is never guessed
    assert h.salesperson == "Charmaine Lane"


def test_missing_header_field_refuses(order):
    h = ax.ExportHeader(**{**KURT_HEADER, "call_letters": " "})
    with pytest.raises(ax.AgencyXmlError, match="call_letters"):
        ax.export_proposal(order, h)


def test_preview_lines_follow_sheet_order(order):
    pv = ax.build_preview(order, ax.ExportHeader(**KURT_HEADER))
    assert [ln.program for ln in pv.lines] == [
        "Cantonese News",
        "Mandarin News",
        "Filipino News/Talk",
        "Vietnamese News /Variety",
        "Cantonese",
        "Mandarin",
        "Filipino",
        "Vietnamese",
    ]
    assert [ln.daypart_name for ln in pv.lines] == [
        "CHINESE",
        "CHINESE",
        "FILIPINO",
        "VIETNAMESE",
        "CHINESE",
        "CHINESE",
        "FILIPINO",
        "VIETNAMESE",
    ]
    assert [ln.rate for ln in pv.lines] == [120.0, 120.0, 100.0, 100.0, 0.0, 0.0, 0.0, 0.0]
    assert [ln.length_sec for ln in pv.lines] == [30, 30, 30, 30, 15, 15, 15, 15]
    assert (pv.flight_start, pv.flight_end) == (date(2026, 7, 27), date(2026, 11, 1))
    assert pv.market_label == "San Francisco Bay Area - Xfinity Channel 3131 and KQTA 15.3"


def test_spots_per_week_option_emits_consolidated_periods(order):
    h = ax.ExportHeader(**KURT_HEADER, include_spots_per_week=True)
    xml_bytes = ax.export_proposal(order, h)
    root = ET.fromstring(xml_bytes)
    ns = {"p": ax.NS_PROPOSAL}
    first = root.find(".//p:AvailLineWithDetailedPeriods", ns)
    periods = first.findall("p:Periods/p:DetailedPeriod", ns)
    assert [
        (p.get("startDate"), p.get("endDate"), p.findtext("p:SpotsPerWeek", namespaces=ns))
        for p in periods
    ] == [("2026-07-27", "2026-10-25", "3"), ("2026-10-26", "2026-11-01", "2")]


@pytest.mark.parametrize(
    "text,expect",
    [
        ("M-F 7p-8p", [("19:00", "20:00", (1, 1, 1, 1, 1, 0, 0))]),
        ("M-Sun  8p-9p", [("20:00", "21:00", (1, 1, 1, 1, 1, 1, 1))]),
        ("M-Sun 11a-1p", [("11:00", "13:00", (1, 1, 1, 1, 1, 1, 1))]),
        ("ROS", [("06:00", "24:00", (1, 1, 1, 1, 1, 1, 1))]),
        ("Sa-Su 7p-12a", [("19:00", "24:00", (0, 0, 0, 0, 0, 1, 1))]),
        ("M-F 11-1p", [("11:00", "13:00", (1, 1, 1, 1, 1, 0, 0))]),
        (
            "M-F 6a-8a; Sat-Sun 7p-11p",
            [("06:00", "08:00", (1, 1, 1, 1, 1, 0, 0)), ("19:00", "23:00", (0, 0, 0, 0, 0, 1, 1))],
        ),
        ("M-F 10:30a-12n", [("10:30", "12:00", (1, 1, 1, 1, 1, 0, 0))]),
    ],
)
def test_parse_daypart(text, expect):
    assert [(d.start, d.end, d.days) for d in ax.parse_daypart(text)] == expect


@pytest.mark.parametrize("text", ["", "M-F", "Tuesdays 7p-8p", "M-F 7p", "M-F 25p-26p", "M-F ???"])
def test_unreadable_daypart_refuses(text):
    with pytest.raises(ax.AgencyXmlError):
        ax.parse_daypart(text)


def test_language_family():
    assert ax.language_family("Cantonese News") == "CHINESE"
    assert ax.language_family("Tagalog Variety") == "FILIPINO"
    assert ax.language_family("Punjabi") == "SOUTH ASIAN"
    assert ax.language_family("Korean Drama") == "KOREAN"


# ── Tampered workbooks must refuse, never export a wrong number ──────────────


def _tampered(tmp_path, mutate):
    """Copy the workbook with its cached VALUES (openpyxl would otherwise save the
    formulas with no cached result, blanking every total) and mutate one cell."""
    dst = tmp_path / "proposal.xlsx"
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    mutate(wb["BAAQMD"])
    wb.save(dst)
    return str(dst)


def test_rate_that_disagrees_with_the_contract_amount_refuses(tmp_path):
    path = _tampered(tmp_path, lambda ws: ws.__setitem__("F20", 150))  # Cantonese News 120 → 150
    with pytest.raises(ValueError, match="Proposed Contract Amount"):
        parse_crispin_xlsx(path)


def test_dropped_week_cell_refuses(tmp_path):
    path = _tampered(tmp_path, lambda ws: ws.__setitem__("H20", None))
    with pytest.raises(ValueError):
        parse_crispin_xlsx(path)


def test_line_with_no_spots_refuses(tmp_path):
    def zero_row(ws):
        for col in "HIJKLMNOPQRSTU":
            ws[f"{col}24"] = 0
        ws["V24"] = 0
        ws["V29"] = 164 - 41

    order = parse_crispin_xlsx(_tampered(tmp_path, zero_row))
    with pytest.raises(ax.AgencyXmlError, match="no spots"):
        ax.build_preview(order, ax.ExportHeader(**KURT_HEADER))


def test_untouched_copy_still_parses(tmp_path):
    order = parse_crispin_xlsx(_tampered(tmp_path, lambda ws: None))
    assert len(order.lines) == 8


# ── Header memory + a two-tab workbook (Winter 2026, gross rates) ──────────────

WINTER = FIX / "proposal_baaqmd_2026_winter.xlsm"


def test_default_header_ignores_a_missing_memory_file(order, tmp_path):
    h = ax.default_header(order, memory_path=tmp_path / "none.json")
    assert (h.buyer_company, h.call_letters) == ("Crispin LLC", "")


def test_remembered_buyer_and_station_code_prefill_the_next_proposal(order, tmp_path):
    mem = tmp_path / "headers.json"
    ax.remember_header(ax.ExportHeader(**KURT_HEADER), path=mem)
    h = ax.default_header(order, memory_path=mem)
    assert h.buyer_company == "Allison Worldwide"
    assert h.call_letters == "3131CA"
    assert h.buyer_name == "Alexander Boyle"  # the sheet's contact still wins when present
    assert (
        ax.remembered_header("bay area quality management district", mem)["call_letters"]
        == "3131CA"
    )


def test_winter_workbook_has_two_proposal_tabs():
    from browser_automation.parsers.crispin_parser import proposal_sheet_names

    assert proposal_sheet_names(str(WINTER)) == ["new creative", "REUSE"]
    new = parse_crispin_xlsx(str(WINTER), "new creative")
    reuse = parse_crispin_xlsx(str(WINTER), "REUSE")
    assert (new.subtitle, reuse.subtitle) == ("NEW CREATIVE", "RE-USE CREATIVE")
    # gross rates reconcile against the 'GROSS Proposed Contract Amount' column
    assert [ln.amount_stated for ln in new.lines][:4] == [6353.1, 6353.1, 3529.5, 5294.25]
    assert [ln.total_spots for ln in new.paid_lines] == [45, 45, 30, 45]
    assert [ln.total_spots for ln in reuse.paid_lines] == [45, 45, 45, 45]


def test_winter_defaults_carry_the_tab_variant(tmp_path):
    new = parse_crispin_xlsx(str(WINTER), "new creative")
    h = ax.default_header(new, memory_path=tmp_path / "none.json")
    assert h.proposal_id == "BAAQMD-2026-NEW-CREATIVE"
    assert h.proposal_name == "Crossings TV - BAAQMD 2026 New Creative"
    assert h.product == "BAAQMD 2026"
    assert h.buyer_company == "CRISPIN"
    pv = ax.build_preview(new, ax.ExportHeader(**{**KURT_HEADER, "proposal_id": h.proposal_id}))
    assert (pv.flight_start, pv.flight_end) == (date(2026, 11, 9), date(2027, 2, 21))
    assert [ln.rate for ln in pv.lines][:4] == [141.18, 141.18, 117.65, 117.65]
    assert ax.validate_proposal_xml(ax.build_proposal_xml(pv)) == []


# ── The page ────────────────────────────────────────────────────────────────


@pytest.fixture
def client(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.templating import Jinja2Templates
    from fastapi.testclient import TestClient

    from web.routes import agency_xml as route

    monkeypatch.setattr(ax, "HEADER_MEMORY_PATH", tmp_path / "headers.json")
    monkeypatch.setattr(
        route, "remember_header", lambda h: ax.remember_header(h, tmp_path / "headers.json")
    )
    monkeypatch.setattr(
        route, "default_header", lambda o: ax.default_header(o, tmp_path / "headers.json")
    )
    app = FastAPI()
    app.include_router(
        route.build_agency_xml_router(Jinja2Templates(directory="src/web/templates"))
    )
    return TestClient(app)


def _shared(**over):
    form = {
        k: (v.isoformat() if isinstance(v, date) else v)
        for k, v in KURT_HEADER.items()
        if k in ("advertiser", "buyer_company", "buyer_name", "call_letters", "send_date")
    }
    form["salesperson"] = "Charmaine Lane"
    form.update(over)
    return form


def test_page_preview_and_download(client):
    assert "Agency XML Export" in client.get("/orders/agency-xml").text
    with WORKBOOK.open("rb") as fh:
        r = client.post("/orders/agency-xml/preview", files={"workbook": (WORKBOOK.name, fh)})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["shared"]["call_letters"] == ""
    assert [s["sheet"] for s in body["sheets"]] == ["BAAQMD"]
    sh = body["sheets"][0]
    assert sh["proposal_id"] == "BAAQMD-2026"
    assert len(sh["lines"]) == 8
    assert sh["lines"][4]["windows"] == [{"start": "06:00", "end": "24:00", "days": "MTWTFSS"}]
    assert sh["paid_total"] == 17820.0

    picked = [
        {
            "sheet": "BAAQMD",
            "proposal_id": "BAAQMD-2026-R1C",
            "proposal_name": "Crossings TV - BAAQMD 2026",
            "product": "BAAQMD 2026",
        }
    ]
    with WORKBOOK.open("rb") as fh:
        r = client.post(
            "/orders/agency-xml/generate",
            data={**_shared(), "sheets": json.dumps(picked)},
            files={"workbook": (WORKBOOK.name, fh)},
        )
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"] == 'attachment; filename="BAAQMD-2026-R1C.xml"'
    assert _tree(ET.fromstring(r.content)) == _tree(ET.parse(ORACLE).getroot())

    # the buyer/station values are remembered for the advertiser's next proposal
    with WINTER.open("rb") as fh:
        r = client.post("/orders/agency-xml/preview", files={"workbook": (WINTER.name, fh)})
    assert r.json()["shared"]["call_letters"] == "3131CA"
    assert r.json()["shared"]["buyer_company"] == "Allison Worldwide"


def test_two_ticked_tabs_download_a_zip(client):
    picked = [
        {
            "sheet": "new creative",
            "proposal_id": "BAAQMD-2026-NEW-CREATIVE",
            "proposal_name": "Crossings TV - BAAQMD 2026 New Creative",
            "product": "BAAQMD 2026",
        },
        {
            "sheet": "REUSE",
            "proposal_id": "BAAQMD-2026-RE-USE-CREATIVE",
            "proposal_name": "Crossings TV - BAAQMD 2026 Re-Use Creative",
            "product": "BAAQMD 2026",
        },
    ]
    with WINTER.open("rb") as fh:
        r = client.post(
            "/orders/agency-xml/generate",
            data={**_shared(), "sheets": json.dumps(picked)},
            files={"workbook": (WINTER.name, fh)},
        )
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"] == 'attachment; filename="BAAQMD-2026-proposals.zip"'
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        assert zf.namelist() == ["BAAQMD-2026-NEW-CREATIVE.xml", "BAAQMD-2026-RE-USE-CREATIVE.xml"]
        for name in zf.namelist():
            assert ax.validate_proposal_xml(zf.read(name)) == []
        root = ET.fromstring(zf.read("BAAQMD-2026-RE-USE-CREATIVE.xml"))
    ns = {"p": ax.NS_PROPOSAL}
    assert root.find("p:Proposal", ns).get("uniqueIdentifier") == "BAAQMD-2026-RE-USE-CREATIVE"
    assert len(root.findall(".//p:AvailLineWithDetailedPeriods", ns)) == 8


def test_duplicate_proposal_ids_refuse(client):
    picked = [
        {"sheet": "new creative", "proposal_id": "X", "proposal_name": "n", "product": "p"},
        {"sheet": "REUSE", "proposal_id": "X", "proposal_name": "n", "product": "p"},
    ]
    with WINTER.open("rb") as fh:
        r = client.post(
            "/orders/agency-xml/generate",
            data={**_shared(), "sheets": json.dumps(picked)},
            files={"workbook": (WINTER.name, fh)},
        )
    assert r.status_code == 400 and "Proposal ID" in r.json()["detail"]


def test_page_refuses_a_blank_station_code(client):
    picked = [
        {"sheet": "BAAQMD", "proposal_id": "BAAQMD-2026-R1C", "proposal_name": "n", "product": "p"}
    ]
    with WORKBOOK.open("rb") as fh:
        r = client.post(
            "/orders/agency-xml/generate",
            data={**_shared(call_letters=""), "sheets": json.dumps(picked)},
            files={"workbook": (WORKBOOK.name, fh)},
        )
    assert r.status_code == 400
    assert "call_letters" in r.json()["detail"]


def test_page_refuses_a_non_workbook(client):
    r = client.post("/orders/agency-xml/preview", files={"workbook": ("io.pdf", b"%PDF-1.4")})
    assert r.status_code == 400
