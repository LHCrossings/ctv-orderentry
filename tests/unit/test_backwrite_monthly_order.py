"""Monthly totals list months in CHRONOLOGICAL order — an order that crosses a
calendar year shows December, then January (Lee, Vida PSE 3173, 2026-10-12).

Before the fix the Sales Confirmation breakdown keyed months by NAME and sorted
by calendar position (January above December), and the Sheet1 pivot sorted
"Month Year" strings alphabetically (August above July)."""

import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

_root = Path(__file__).resolve().parents[2]
for _p in [str(_root), str(_root / "src")]:
    if _p not in sys.path:
        sys.path.insert(0, _p)

from backwrite.transformer import (  # noqa: E402
    _fill_monthly_breakdown,
    _fill_pivot,
    _month_label,
    _month_sort_key,
)


def _sc_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["", "line 1"])
    ws.append(["Month", "Gross", "Net"])
    ws.append(["July", 1, 1])
    ws.append(["August", 1, 1])
    ws.append(["Total", 2, 2])
    return ws


def _month_column(ws, start_row):
    out = []
    for r in range(start_row, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None or str(v).lower().startswith("total"):
            break
        out.append(v)
    return out


def test_sales_confirmation_breakdown_is_chronological_across_a_year_boundary():
    ws = _sc_sheet()
    gross = {date(2027, 1, 1): 600.0, date(2026, 12, 1): 6400.0}
    _fill_monthly_breakdown(ws, gross, {}, last_line_row=1, agency_fee=0.0)
    assert _month_column(ws, 3) == ["December", "January"]
    assert ws.cell(3, 2).value == 6400.0 and ws.cell(4, 2).value == 600.0


def test_sales_confirmation_breakdown_keeps_calendar_order_within_a_year():
    ws = _sc_sheet()
    gross = {date(2026, 11, 1): 1.0, date(2026, 9, 1): 2.0, date(2026, 10, 1): 3.0}
    _fill_monthly_breakdown(ws, gross, {}, last_line_row=1, agency_fee=0.0)
    assert _month_column(ws, 3) == ["September", "October", "November"]


def test_legacy_month_name_keys_still_sort_by_calendar():
    assert sorted(["March", "January", "February"], key=_month_sort_key) == [
        "January",
        "February",
        "March",
    ]
    assert _month_label("March") == "March"
    assert _month_label(date(2026, 12, 1)) == "December"


def test_pivot_sorts_by_date_not_by_name():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["<month>", "Gross", "Net"])
    ws.append(["Grand Total", 0, 0])
    rows = [
        {"month": datetime(2027, 1, 1), "gross_rate": 1.0, "station_net": 1.0},
        {"month": datetime(2026, 8, 1), "gross_rate": 2.0, "station_net": 2.0},
        {"month": datetime(2026, 12, 1), "gross_rate": 3.0, "station_net": 3.0},
        {"month": datetime(2026, 7, 1), "gross_rate": 4.0, "station_net": 4.0},
    ]
    _fill_pivot(ws, rows)
    assert [ws.cell(r, 1).value for r in range(1, 5)] == [
        "July 2026",
        "August 2026",
        "December 2026",
        "January 2027",
    ]
    assert [ws.cell(r, 2).value for r in range(1, 5)] == [4.0, 2.0, 3.0, 1.0]
    assert ws.cell(5, 1).value == "Grand Total" and ws.cell(5, 2).value == 10.0


def _sc_sheet_one_row():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["", "line 1"])
    ws.append(["Month", "Gross", "Net"])
    ws.append(["July", 1, 1])
    ws.append(["Total", 1, 1])
    return ws


def test_calendar_rows_carrying_air_dates_collapse_to_one_row_per_month():
    """EtereBridge's Month column is the raw AIR DATE on Calendar contracts; the
    breakdown must still show one row per month (RWNY, Lee 2026-10-24: 70 rows
    reading 'October $500 / October $600 / ...')."""
    from backwrite.transformer import _month_key, generate_excel  # noqa: F401

    rows = [
        {"month": datetime(2026, 10, 27), "gross_rate": 500.0, "station_net": 500.0},
        {"month": datetime(2026, 10, 28), "gross_rate": 600.0, "station_net": 600.0},
        {"month": datetime(2026, 11, 2), "gross_rate": 100.0, "station_net": 100.0},
        {"month": datetime(2026, 11, 3), "gross_rate": 0.0, "station_net": 0.0},
        {"month": datetime(2026, 12, 30), "gross_rate": 200.0, "station_net": 200.0},
    ]
    assert _month_key(datetime(2026, 10, 27)) == date(2026, 10, 1)
    assert _month_key(date(2026, 10, 1)) == date(2026, 10, 1)

    # Pivot (Sheet1) path
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["<month>", "Gross", "Net"])
    ws.append(["Grand Total", 0, 0])
    _fill_pivot(ws, rows)
    assert [ws.cell(r, 1).value for r in range(1, 4)] == [
        "October 2026",
        "November 2026",
        "December 2026",
    ]
    assert [ws.cell(r, 2).value for r in range(1, 4)] == [1100.0, 100.0, 200.0]
    assert ws.cell(4, 1).value == "Grand Total"

    # Sales Confirmation breakdown path: aggregate the way generate_excel does
    from collections import defaultdict

    mg: dict = defaultdict(float)
    for rr in rows:
        mg[_month_key(rr["month"])] += rr["gross_rate"]
    assert set(mg) == {date(2026, 10, 1), date(2026, 11, 1), date(2026, 12, 1)}
    sc = _sc_sheet_one_row()
    _fill_monthly_breakdown(sc, dict(mg), {}, last_line_row=1, agency_fee=0.0)
    assert _month_column(sc, 3) == ["October", "November", "December"]
    assert [sc.cell(r, 2).value for r in (3, 4, 5)] == [1100.0, 100.0, 200.0]
