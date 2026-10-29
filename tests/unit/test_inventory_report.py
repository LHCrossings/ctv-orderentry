"""Inventory (avails) report rules (2026-10-29): open COMS time per program and breaks with room."""

from business_logic.services.inventory_report import (
    DAY_FRAMES,
    HOUR_FRAMES,
    build,
    build_excel,
    day_pattern,
    frames_to_bcast,
    parse_block_name,
)

H = HOUR_FRAMES
F30 = 899  # :30 in frames at 29.97 (899.1)
F15 = 450


def _seg(date, bstart, seg_rel, dur=F30 * 4, market=1, name="P - To the Point REV", bid=10, bdur=H):
    return {
        "date": date,
        "market": market,
        "block_id": bid,
        "name": name,
        "bstart": bstart,
        "bdur": bdur,
        "seg_off": bstart + seg_rel,
        "seg_dur": dur,
        "seg_id": seg_rel,
    }


def _spot(date, offset, dur=F30, market=1, ora=None):
    return {
        "date": date,
        "market": market,
        "offset": offset,
        "ora": ora if ora is not None else offset,
        "dur": dur,
    }


def _to_the_point(date):
    # 2p-3p, four 2:00 breaks at :14, :25, :40, :55
    return [_seg(date, 14 * H, r) for r in (14 * 60 * 30, 25 * 60 * 30, 40 * 60 * 30, 55 * 60 * 30)]


def test_block_name_prefix_is_the_language():
    assert parse_block_name("P - To the Point REV") == ("P", "To the Point REV")
    assert parse_block_name("K- Korean Drama Saturday") == ("K", "Korean Drama Saturday")
    assert parse_block_name("SA - Big News REV")[0] == "SA"
    assert parse_block_name("Hm - Hmong Report")[0] == "Hm"
    assert parse_block_name("Shop LC 60") == ("?", "Shop LC 60")


def test_broadcast_time_formatting_wraps_past_midnight():
    assert frames_to_bcast(14 * H) == "2p"
    assert frames_to_bcast(int(15.5 * H)) == "3:30p"
    assert frames_to_bcast(12 * H) == "12p"
    assert frames_to_bcast(25 * H) == "1a"
    assert frames_to_bcast(6 * H) == "6a"


def test_day_patterns_match_maijas_sheet():
    assert day_pattern({0, 1, 2, 3, 4}) == "M-F"
    assert day_pattern({5}) == "Sa"
    assert day_pattern({6}) == "Su"
    assert day_pattern({5, 6}) == "Sa-Su"
    assert day_pattern(set(range(7))) == "M-Su"
    assert day_pattern({0, 2, 4}) == "M,W,F"


def test_open_time_and_breaks_with_room():
    """Four 2:00 breaks (480 s): sell :30 in break 1, fill break 2 completely, nothing else ->
    open 480-30-120 = 330 s, breaks with room 3 of 4."""
    segs = _to_the_point("2026-11-03")
    sold = [_spot("2026-11-03", segs[0]["seg_off"])] + [
        _spot("2026-11-03", segs[1]["seg_off"]) for _ in range(4)
    ]
    rep = build(segs, sold)
    row = rep["markets"][0]["groups"][0]["rows"][0]
    assert (row["lang"], row["title"], row["time"]) == ("P", "To the Point REV", "2p-3p")
    assert row["cap_s"] == 480 and row["sold_s"] == 150 and row["open_s"] == 330
    assert row["breaks_per_airing"] == 4 and row["breaks_open"] == 3 and row["breaks"] == 4
    assert rep["totals"]["unattributed_spots"] == 0


def test_a_break_with_less_than_15s_left_has_no_room():
    segs = [_seg("2026-11-03", 14 * H, 1000, dur=F30)]  # one :30 break
    rep = build(segs, [_spot("2026-11-03", 14 * H + 1000, dur=F15 + 60)])  # :17 sold, 13 s left
    row = rep["markets"][0]["groups"][0]["rows"][0]
    assert row["open_s"] == 13 and row["breaks_open"] == 0


def test_spot_booked_at_an_offset_no_break_has_goes_to_the_nearest_break_of_its_block():
    segs = _to_the_point("2026-11-03")
    odd = _spot("2026-11-03", 14 * H + 999, ora=14 * H + 24 * 60 * 30)  # ORA near the :25 break
    rep = build(segs, [odd])
    row = rep["markets"][0]["groups"][0]["rows"][0]
    assert row["sold_s"] == 30 and rep["totals"]["unattributed_spots"] == 0
    assert row["days"]["2026-11-03"]["breaks_open"] == 4  # :30 in a 2:00 break still has room


def test_language_and_time_filters():
    segs = _to_the_point("2026-11-03") + [
        _seg("2026-11-03", 19 * H, 1000, name="C - Cantonese Primetime News", bid=20)
    ]
    assert [g["lang"] for g in build(segs, [])["markets"][0]["groups"]] == ["P", "C"]
    assert [g["lang"] for g in build(segs, [], languages=["C"])["markets"][0]["groups"]] == ["C"]
    assert [
        g["lang"] for g in build(segs, [], t_from="18:00", t_to="23:59")["markets"][0]["groups"]
    ] == ["C"]
    # post-midnight tail: a 1a block is at 25h and a 00:00-05:59 window must find it
    late = [_seg("2026-11-03", 25 * H, 1000, name="M - Late Movie", bid=30)]
    assert build(late, [], t_from="00:00", t_to="05:59")["markets"][0]["groups"][0]["lang"] == "M"
    assert build(late, [], t_from="06:00", t_to="23:59")["markets"] == []


def test_rows_aggregate_over_the_range_with_the_day_pattern():
    segs = []
    for d in ("2026-11-02", "2026-11-03", "2026-11-04", "2026-11-05", "2026-11-06"):  # Mon-Fri
        segs += _to_the_point(d)
    segs += [_seg("2026-11-07", 16 * H, 1000, name="T - Filipino Drama", bid=40)]  # Saturday
    rep = build(segs, [_spot("2026-11-03", segs[0]["seg_off"])])
    rows = {
        (r["lang"], r["day_pattern"]): r
        for m in rep["markets"]
        for g in m["groups"]
        for r in g["rows"]
    }
    p = rows[("P", "M-F")]
    assert p["airings"] == 5 and p["cap_s"] == 2400 and p["open_s"] == 2370
    assert p["open_per_airing_s"] == 474 and p["breaks_open"] == 20
    assert rows[("T", "Sa")]["airings"] == 1
    assert rep["dates"][0] == "2026-11-02" and len(rep["dates"]) == 6


def test_markets_are_separate_tables():
    segs = _to_the_point("2026-11-03") + [_seg("2026-11-03", 14 * H, 1000, market=4, bid=50)]
    rep = build(segs, [_spot("2026-11-03", 14 * H + 1000, market=4)])
    assert [m["market_code"] for m in rep["markets"]] == ["NYC", "SFO"]
    assert rep["markets"][1]["sold_s"] == 30 and rep["markets"][0]["sold_s"] == 0


def test_excel_has_a_summary_and_one_daily_sheet_per_market():
    import io

    import openpyxl

    segs = _to_the_point("2026-11-03") + [_seg("2026-11-03", 14 * H, 1000, market=4, bid=50)]
    rep = build(segs, [])
    rep.update({"date_from": "2026-11-03", "date_to": "2026-11-03", "generated_at": "x"})
    wb = openpyxl.load_workbook(io.BytesIO(build_excel(rep)))
    assert wb.sheetnames == ["Inventory", "NYC daily", "SFO daily"]
    ws = wb["Inventory"]
    assert ws["A4"].value == "Market" and ws["A5"].value == "NYC" and ws["I5"].value == 480


def test_frames_constants():
    assert DAY_FRAMES == 24 * HOUR_FRAMES
