"""Compile Weekly Logs: a corrupt upload is named (Lee 9/28: the page showed only
"Bad CRC-32 for file 'xl/workbook.xml'" with no way to tell which of the 11
workbooks it came from)."""

import io
import sys
from pathlib import Path

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.templating import Jinja2Templates
from fastapi.testclient import TestClient

_root = Path(__file__).resolve().parents[2]
for p in [str(_root), str(_root / "src")]:
    if p not in sys.path:
        sys.path.insert(0, p)

from orchestration.config import ApplicationConfig  # noqa: E402
from web.routes.orders import build_router  # noqa: E402


def _book(tab: str, rows=()) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tab
    ws.append(["header"])
    for r in rows:
        ws.append(list(r))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    d = tmp_path_factory.mktemp("cfg")
    cfg = ApplicationConfig(
        incoming_dir=d / "in",
        processed_dir=d / "out",
        error_dir=d / "err",
        customer_db_path=d / "c.db",
    )
    app = FastAPI()
    app.include_router(build_router(cfg, Jinja2Templates(directory=str(d))))
    return TestClient(app)


def _post(client, book: bytes, logs: dict[str, bytes]):
    files = [("billing_book", ("Master Billing Sheet 2609.xlsm", book))]
    files += [("log_files", (name, data)) for name, data in logs.items()]
    return client.post("/api/billing/compile-logs/aggregate", files=files)


def test_good_uploads_compile(client):
    logs = {
        "NYC Log - 092126.xlsm": _book("Master for Billing", [("a", 1)]),
        "CMP Log - 092126.xlsm": _book("Master for Billing", [("b", 2)]),
    }
    r = _post(client, _book("Master"), logs)
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert [tuple(x) for x in wb["Master"].iter_rows(values_only=True)] == [
        ("header", None),  # appended rows widen the sheet to two columns
        ("a", 1),
        ("b", 2),
    ]


def test_truncated_log_is_named(client):
    good = _book("Master for Billing", [("a", 1)])
    r = _post(client, _book("Master"), {"SEA Log - 092126.xlsm": good[: len(good) // 2]})
    assert r.status_code == 400
    assert r.json()["detail"].startswith(
        "Log file SEA Log - 092126.xlsm is not a readable workbook"
    )
    assert "truncated or still copying" in r.json()["detail"]


def test_truncated_billing_book_is_named(client):
    book = _book("Master")
    r = _post(client, book[:-40], {"SEA Log - 092126.xlsm": _book("Master for Billing")})
    assert r.status_code == 400
    assert r.json()["detail"].startswith("Billing book is not a readable workbook")
